from xml.etree import ElementTree

import struct
from django.db.models import Q
from django.core.validators import MinValueValidator
from django.db import models
from django.utils.translation import ugettext_lazy as _

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.writer.excel import save_virtual_workbook

import logging

logger = logging.getLogger(__name__)


def str_to_ord(string):
    string = string.upper()
    value = 0
    for i in range(0, len(string)):
        value += (ord(string[i]) - 64) * pow(26, len(string) - i - 1)
    return value


def get_span(start, end):
    start_row = int(start.lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))
    start_column = str_to_ord(start.rstrip('0123456798'))

    end_row = int(end.lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))
    end_column = str_to_ord(end.rstrip('0123456798'))

    column_span = end_column - start_column
    row_span = end_row - start_row

    row_span = row_span + 1 if row_span else row_span
    column_span = column_span + 1 if column_span else column_span

    return column_span, row_span


class DocumentManager(models.Manager):
    def get_current(self, id):
        instance = self.get_queryset().get(
            (Q(replaces=id) & Q(current=True) & Q(replaces__isnull=False)) |
            (Q(id=id)) & Q(current=True) & Q(replaces__isnull=True))
        return instance

    def all_current(self):
        return self.get_queryset().filter(current=True)


class Document(models.Model):
    OPEN = 1
    REQUEST_ONLY = 2
    LOCKED = 3
    STATUS = ((OPEN, _("Open")),
              (REQUEST_ONLY, _("Semi locked")),
              (LOCKED, _("Locked")))
    file = models.FileField(blank=True, null=True)
    replaces = models.ForeignKey('self', blank=True, null=True, )
    name = models.CharField(max_length=100)
    created = models.DateTimeField(auto_now_add=True)
    current = models.BooleanField(default=True)
    status = models.IntegerField(choices=STATUS, default=OPEN)
    worksheet = models.IntegerField(default=0)

    objects = DocumentManager()

    _theme_colors = []
    _color_set = set()

    def __str__(self):
        return self.name

    def save(self, *args, **kwargs):
        parse_file = False
        if not self.pk:
            if self.current:
                Document.objects.filter(
                    (Q(replaces=self.replaces) & Q(replaces__isnull=False)) |
                    Q(id=self.replaces_id)).update(current=False)
            parse_file = True
        super(Document, self).save(*args, **kwargs)

        if parse_file:
            self.parse_file()

    _workbook = None

    @property
    def workbook(self):
        if not self._workbook:
            self._workbook = load_workbook(self.file.path)
        return self._workbook

    _skip_list = None

    @property
    def skip_list(self):
        if not self._skip_list:
            self.create_skip_and_start_list()
        return self._skip_list

    _start_cells = None

    @property
    def start_cells(self):
        if not self._start_cells:
            self.create_skip_and_start_list()
        return self._start_cells

    def create_skip_and_start_list(self):
        start_cells = dict()
        work_sheet = self.workbook.worksheets[self.worksheet]

        skip_list = work_sheet.merged_cells.copy()
        for cell_range in work_sheet.merged_cell_ranges:
            start, end = cell_range.split(':')
            skip_list.remove(start)

            column_span, row_span = get_span(start, end)
            cell_span = dict()
            if row_span:
                cell_span['row_span'] = row_span
            if column_span:
                cell_span['column_span'] = column_span
            start_cells[start] = cell_span

        self._start_cells = start_cells
        self._skip_list = skip_list

    def get_theme_color(self, cell):
        fg_color = cell.fill.fgColor
        if not fg_color.type == "theme":
            return None

        theme_index = fg_color.value
        if not self._theme_colors:
            self._theme_colors = self.parse_theme_colors()
        try:
            return self._theme_colors[theme_index]
        except IndexError:
            return None

    def create_xlsx(self):
        workbook = self.workbook
        work_sheet = workbook.worksheets[self.worksheet]

        for cell in self.cell_set.all():
            work_sheet[cell.coordinate].value = cell.value
        return save_virtual_workbook(workbook)

    @property
    def url_id(self):
        return self.replaces_id or self.pk

    def parse_theme_colors(self):
        colors = []
        if not self.workbook.loaded_theme:
            return colors
        try:
            name_space = {
                'ns': "http://schemas.openxmlformats.org/drawingml/2006/main"}
            root = ElementTree.fromstring(self.workbook.loaded_theme)
            theme_elements = root.find('ns:themeElements', name_space)
            color_schemes = theme_elements.findall('ns:clrScheme', name_space)
            first_color_scheme = color_schemes[0]

            for c in ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2',
                      'accent3', 'accent4', 'accent5', 'accent6']:
                accent = first_color_scheme.find('ns:' + c, name_space)

                if 'window' in accent.getchildren()[0].attrib['val']:
                    colors.append(accent.getchildren()[0].attrib['lastClr'])
                else:
                    colors.append(accent.getchildren()[0].attrib['val'])
        except AttributeError:
            logger.exception('Parsing theme colors failed')
        return colors

    @staticmethod
    def get_name_for_color(color):
        return ''.join(['color_', color])

    def save_colors(self):
        for color in self._color_set:
            name = Document.get_name_for_color(color)
            DocumentColors.objects.create(document=self,
                                          color=color,
                                          name=name)

    def get_color_from_cell(self, cell):
        color = cell.fill.fgColor
        color_value = None

        # skip wrong black values
        if color.value == '00000000':
            color_value = None
        elif color.type == 'rgb':
            color_value = color.value
        elif color.type == 'indexed':
            color_value = ''.join(['FF', COLOR_INDEX[color.value][2:]])
        elif color.type == 'theme':
            theme_color = self.get_theme_color(cell)
            tint = "%02X" % int(color.tint * 255)
            color_value = ''.join([tint, theme_color])

        if not color_value:
            return ""

        name = Document.get_name_for_color(color_value)
        self._color_set.add(color_value)
        return name

    def prepare_cell(self, cell, is_first_cell):
        color_name = self.get_color_from_cell(cell)
        row_span = None
        column_span = None

        if cell.coordinate in self.start_cells:
            row_span = self.start_cells[cell.coordinate].get('row_span',  None)
            column_span = self.start_cells[cell.coordinate].get('column_span',
                                                                None)
        db_cell = Cell(coordinate=cell.coordinate,
                       value=cell.value or "",
                       color_name=color_name,
                       row_span=row_span,
                       column_span=column_span,
                       document=self,
                       horizontal_alignment=cell.alignment.horizontal,
                       first_cell=is_first_cell, )
        return db_cell

    def parse_file(self):
        work_sheet = self.workbook.worksheets[self.worksheet]

        cell_list = list()
        for row in work_sheet.rows:
            # TODO: Check if all columns have same length
            # TODO: skip empty cells/rows. E.g. skip row if first X(=10?)
            #       columns are empty

            is_first_cell = True
            for cell in row:
                if cell.coordinate in self.skip_list:
                    continue
                db_cell = self.prepare_cell(cell, is_first_cell)
                if db_cell:
                    cell_list.append(db_cell)

                is_first_cell = False
                logger.debug("created cell %s" % cell.coordinate)

        cell_list[len(cell_list) - 1].last_cell = True
        Cell.objects.bulk_create(cell_list)
        self.save_colors()


class Cell(models.Model):
    coordinate = models.CharField(max_length=15)
    value = models.CharField(max_length=255, default="", blank=True)
    color_name = models.CharField(max_length=20, blank=True)
    row_span = models.IntegerField(blank=True,
                                   null=True,
                                   validators=MinValueValidator(1))
    column_span = models.IntegerField(blank=True,
                                      null=True,
                                      validators=MinValueValidator(1))
    document = models.ForeignKey(Document)
    horizontal_alignment = models.CharField(max_length=20,
                                            null=True,
                                            blank=True)
    first_cell = models.BooleanField(default=False)
    last_cell = models.BooleanField(default=False)

    class Meta:
        ordering = ['id', ]

    def __str__(self):
        return '%s (Document: %d)' % (self.coordinate, self.document_id)

    @property
    def attributes(self):
        attributes = str()
        if self.row_span:
            attributes = 'rowspan=%d' % self.row_span
        if self.column_span:
            attributes = ''.join([attributes, 'colspan=%d' % self.column_span])
        return attributes

    @property
    def class_tags(self):
        classes = self.color_name
        if self.horizontal_alignment:
            classes = ''.join([classes, ' ', self.horizontal_alignment])
        return classes


class DocumentColors(models.Model):
    name = models.CharField(max_length=20, db_index=True)
    color = models.CharField(max_length=8)
    document = models.ForeignKey(Document)

    class Meta:
        unique_together = (('name', 'document'),)

    def __str__(self):
        return "%s (%s)" % (self.name, self.color)

    @staticmethod
    def hex_to_rgba(value):
        value = value.lstrip('#')
        lv = len(value)
        argb = [int(value[i:i + lv // 4], 16) for i in range(0, lv, lv // 4)]
        rgba = 'rgba(%d,%d,%d,%.2f)' % \
               (argb[1], argb[2], argb[3], argb[0] / 255)
        return rgba

    @property
    def css_color(self):
        if len(self.color) == 6:
            return ''.join(['#', self.color])
        return DocumentColors.hex_to_rgba(self.color)
